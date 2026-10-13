#!/usr/bin/env python3
"""Build the CME slide deck (.pptx) and slide-by-slide narration (.docx)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from docx import Document
from docx.shared import Pt as DPt, RGBColor as DRGB
from slides_part1 import S1
from slides_part2 import S2
from slides_part3 import S3
from slides_part4 import S4

SLIDES = S1 + S2 + S3 + S4
print(f"Total slides: {len(SLIDES)}")

# Palette: deep plum, soft gold, ivory
PLUM = RGBColor(0x3B, 0x1F, 0x3E)
GOLD = RGBColor(0xC9, 0xA2, 0x5E)
IVORY = RGBColor(0xFA, 0xF6, 0xF0)
GREY = RGBColor(0x5A, 0x52, 0x5E)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def add_bg(slide, color):
    from pptx.oxml.ns import qn
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color

def add_rect(slide, l, t, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh

for i, s in enumerate(SLIDES):
    slide = prs.slides.add_slide(blank)
    is_title = (i == 0)
    is_section = s["t"].startswith("MODULE") or s["t"].startswith("PRETEST") or s["t"].startswith("POSTTEST")
    if is_title:
        add_bg(slide, PLUM)
        add_rect(slide, Inches(0), Inches(6.9), prs.slide_width, Inches(0.12), GOLD)
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(3.4))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; r = p.add_run(); r.text = s["t"]
        r.font.size = Pt(34); r.font.bold = True; r.font.color.rgb = IVORY
        p.alignment = PP_ALIGN.CENTER
        for j, line in enumerate(s["b"]):
            p = tf.add_paragraph(); r = p.add_run(); r.text = line
            r.font.size = Pt(18 if j < 2 else 15)
            r.font.color.rgb = GOLD if j < 2 else IVORY
            p.alignment = PP_ALIGN.CENTER; p.space_before = Pt(8)
    else:
        add_bg(slide, IVORY)
        add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.15), PLUM)
        add_rect(slide, Inches(0), Inches(1.15), prs.slide_width, Inches(0.05), GOLD)
        tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.22), Inches(12.1), Inches(0.85))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; r = p.add_run(); r.text = s["t"]
        r.font.size = Pt(26 if len(s["t"]) < 60 else 21); r.font.bold = True; r.font.color.rgb = IVORY
        body = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4))
        btf = body.text_frame; btf.word_wrap = True
        for j, line in enumerate(s["b"]):
            p = btf.paragraphs[0] if j == 0 else btf.add_paragraph()
            r = p.add_run()
            quiz = line.startswith("Q")
            r.text = ("▸  " if not quiz else "") + line
            r.font.size = Pt(16 if quiz else 19)
            r.font.color.rgb = GREY if quiz else PLUM
            p.space_after = Pt(10 if quiz else 16)
        # footer
        ft = slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.4))
        fp = ft.text_frame.paragraphs[0]; fr = fp.add_run()
        fr.text = "© 2026 Dr. Amr Ismail, MD | 3D Rejuvenation Academy — Approved CPD Provider | www.3drejuvenationcode.com"
        fr.font.size = Pt(9); fr.font.color.rgb = GREY
        pn = ft.text_frame.add_paragraph(); pr_ = pn.add_run(); pr_.text = f"{i+1} / {len(SLIDES)}"
        pr_.font.size = Pt(9); pr_.font.color.rgb = GOLD

prs.save("/workspace/output/Evidence_Based_Advanced_Facial_Aesthetics_Injectables_2026.pptx")
print("PPTX saved")

# ---- Narration DOCX ----
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = DPt(11)
t = doc.add_heading("Evidence-Based Advanced Facial Aesthetics Injectables in Skin & Hair Longevity", 0)
doc.add_paragraph("THE DERMAL EXTRACELLULAR MATRIX — CHEMICAL INTELLIGENCE").runs[0].bold = True
doc.add_paragraph("Exosomes • Polynucleotides • Biomimetic Peptides | 2026 Literature Review")
doc.add_paragraph("Slide-by-Slide Narration Script")
doc.add_paragraph("By Dr. Amr Ismail, MD | Founder, 3D Rejuvenation Academy | © 2026 All rights reserved")
doc.add_paragraph("www.3drejuvenationcode.com | 3D Rejuvenation Academy — Approved CPD Provider")
doc.add_paragraph("")
for i, s in enumerate(SLIDES):
    h = doc.add_heading(f"Slide {i+1}: {s['t']}", level=1)
    for r in h.runs: r.font.color.rgb = DRGB(0x3B, 0x1F, 0x3E)
    p = doc.add_paragraph()
    run = p.add_run("On-screen bullets: "); run.bold = True
    p.add_run(" | ".join(s["b"]))
    p2 = doc.add_paragraph()
    run2 = p2.add_run("Narration: "); run2.bold = True
    p2.add_run(s["n"])
doc.save("/workspace/output/Slide_by_Slide_Narration.docx")
print("Narration DOCX saved")
