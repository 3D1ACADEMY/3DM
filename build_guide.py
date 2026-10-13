#!/usr/bin/env python3
"""Build the Expanded Practitioner Pocket Guide (.docx)."""
from docx import Document
from docx.shared import Pt, RGBColor

PLUM = RGBColor(0x3B, 0x1F, 0x3E)
doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)

def H(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs: r.font.color.rgb = PLUM

def P(text, bold=False, italic=False):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic

def B(items):
    for it in items: doc.add_paragraph(it, style="List Bullet")

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 4"
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = htxt
        for r in c.paragraphs[0].runs: r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = v
    doc.add_paragraph()

doc.add_heading("Expanded Practitioner Pocket Guide", 0)
doc.add_paragraph("Evidence-Based Advanced Facial Aesthetics Injectables in Skin & Hair Longevity").runs[0].bold = True
P("THE DERMAL EXTRACELLULAR MATRIX — CHEMICAL INTELLIGENCE")
P("Exosomes • Polynucleotides • Biomimetic Peptides | 2026 Literature Review")
P("Companion to the 77-slide CME activity (10-item pretest & posttest included).")
P("By Dr. Amr Ismail, MD | Founder, 3D Rejuvenation Academy")
P("www.3drejuvenationcode.com | 3D Rejuvenation Academy — Approved CPD Provider")
P("© 2026 Dr. Amr Ismail, MD. All rights reserved.", italic=True)
P("Disclaimer: educational material; verify local regulation and product licensure. Injectable exosomes are investigational/unapproved in most jurisdictions.", italic=True)
doc.add_page_break()

H("1. How to Use This Guide — Evidence Grades")
P("Every recommendation carries an evidence grade. Counsel patients with these grades, not with marketing claims.")
table(["Grade","Meaning","Practice implication"], [
 ["A","Multiple RCTs / meta-analyses","Practice-changing; may be presented as established"],
 ["B","Single RCT or consistent split-face trials","Supportive; reasonable to offer with informed consent"],
 ["C","Observational studies / case series","Hypothesis-generating; frame as emerging"],
 ["D","Preclinical / mechanistic only","Investigational; never present as proven"],
])

H("2. The Dermal ECM in 60 Seconds")
B([
 "Collagen I (80–85%) + III (10–15%) = ~70% of dermal dry weight; tensile strength and youthful architecture.",
 "Elastin/fibrillin fibers give recoil; UV drives solar elastosis.",
 "Hyaluronic acid binds ~1000× its weight in water; dermal half-life < 24 h.",
 "Fibroblasts read matrix integrity via integrins (FAK/ERK, YAP/TAZ): fragmented collagen switches collagen synthesis OFF.",
 "Photoaging: UV → AP-1 → MMP-1/3/9 ↑ + TGF-β signaling ↓ = chronic collagen deficit.",
 "Senescent fibroblasts secrete SASP (IL-6, IL-8, MMPs) — 'inflammaging'.",
 "Hair: dermal papilla ECM volume predicts follicle caliber; AGA shows DP ECM shrinkage + perifollicular fibrosis.",
])
P("Clinical translation: interventions either REPLACE (fillers), SIGNAL (exosomes, polynucleotides, peptides), PROTECT (SPF, antioxidants), or REMOVE (devices, senolytic strategies). Best results combine ≥ 2.", bold=True)

H("3. Exosomes")
table(["Item","Summary"], [
 ["Definition","30–150 nm endosomal extracellular vesicles carrying proteins, lipids, mRNA, miRNA"],
 ["Key markers","CD9, CD63, CD81 (tetraspanins); characterize by NTA, TEM, Western blot"],
 ["Sources","MSC (adipose/bone marrow/umbilical), platelet/PRP, plant-derived (cosmetic only)"],
 ["Mechanism","miRNA cargo → fibroblast activation, MMP suppression, anti-inflammatory, angiogenic, Wnt/β-catenin in DP cells"],
 ["Skin evidence","Level C — small split-face adjuvant studies with microneedling/laser; no large injectable RCT"],
 ["Hair evidence","Level C — small AGA trials report density gains at 12–24 weeks; no head-to-head vs minoxidil"],
 ["Regulatory","NOT approved for injection (FDA warnings issued; EMA: ATMP pathway). Topical cosmetic ≠ injectable approval"],
 ["Safety","Theoretical: immunogenicity, infection, oncogenic cargo. Serious events linked to unregulated injectables. Long-term data absent"],
 ["Practical position","Topical/post-procedure adjuvant where legal; injection only in trials/permissive jurisdictions; explicit investigational consent"],
])
P("Red flags when sourcing: no characterization data, no potency assay, no sterility/endotoxin certification, 'miracle' claims, no registered trial.", bold=True)

H("4. Polynucleotides (PN / PDRN)")
table(["Item","Summary"], [
 ["Definition","Highly purified, depyrogenated DNA polymers (usually salmon/trout-derived); PDRN = drug-grade shorter fraction"],
 ["Mechanism","Adenosine A2A receptor activation (anti-inflammatory, pro-angiogenic, pro-fibroblast) + nucleotide salvage substrate + hydration"],
 ["Skin evidence","Level B — multiple split-face RCTs: elasticity, hydration, roughness, fine lines, dermal density on ultrasound"],
 ["Special zones","Periocular, neck, décolletage, hands — Level B–C; excellent for thin crepey skin"],
 ["Hair evidence","Level C — pilot studies in AGA/telogen effluvium; trichoscopy density and anagen:telogen improvements"],
 ["Typical protocol","2 mL/session intradermal (micro-depot 0.02–0.05 mL or nappage); 3–4 sessions q2–4 weeks; maintenance q3–6 months"],
 ["Safety","Excellent record; transient papules 24–72 h, bruising. Screen fish-protein allergy; avoid in pregnancy/active infection"],
 ["Position","First-line evidence-graded injectable for skin quality in this review"],
])

H("5. Biomimetic Peptides")
table(["Class","Example","Action","Evidence"], [
 ["Signal","Palmitoyl pentapeptide-4 (pal-KTTKS)","Collagen I fragment mimic → ECM synthesis","Level B (topical RCT)"],
 ["Carrier","GHK-Cu","Cu²⁺ delivery; ↑ collagen/elastin/GAG, lysyl oxidase, SOD; falls ~60% age 20→60","Level C"],
 ["Neurotransmitter-inhibiting","Acetyl hexapeptide-8 (argireline)","Topical botox-like SNARE modulation","Level C (topical)"],
 ["Hair peptides","Biotinoyl tripeptide-1, acetyl tetrapeptide-3","Follicle anchoring claims","Level D"],
 ["Injectable boosters","Peptide + non-cross-linked HA cocktails","Skin quality, post-procedure recovery","Level C"],
])
B([
 "Topical use: daily leave-on; expect visible change at 8–12 weeks.",
 "Injectable boosters: 2–3 monthly sessions, then quarterly maintenance.",
 "Copper peptides: avoid excessive local dosing (pro-oxidant at high concentrations).",
 "Always ask for studies on the FINAL formulation, not the raw molecule; vehicle determines penetration.",
])
H("6. Matrix-Centric Protocol Builder")
table(["Presentation","First-line","Adjunct","Backbone"], [
 ["Crepey periocular / fine lines","PN course (3–4 sessions)","Peptide booster maintenance","SPF, retinoid, antioxidant"],
 ["General skin quality / early laxity","PN ± non-cross-linked HA","Hyperdilute CaHA (scaffold) — staged","SPF, retinoid"],
 ["Post-laser / RF recovery","Standard wound care","Topical exosome or peptide adjuvant (where legal)","SPF, barrier repair"],
 ["Early AGA (men)","Finasteride + minoxidil (standard care)","PN or PRP scalp series; exosomes investigational","Trichoscopy q12 wks"],
 ["Telogen effluvium","Treat trigger","PN scalp course; peptide leave-on serum","Ferritin/thyroid/vitamin D workup"],
 ["Perimenopausal skin thinning","PN + HA skin booster","Hormonal review with GP/gynecology","SPF, retinoid, protein intake"],
])
P("Sequencing rule: energy device first → regenerative injectable after re-epithelialization (or per validated protocol). Do not stack redundant signalers in one session.", bold=True)

H("7. Safety Governance Checklist")
B([
 "Verify source, lot, and QC documentation for every biologic (CoA, sterility, endotoxin, characterization).",
 "Consent form distinguishes evidence grades; investigational status stated explicitly (exosomes).",
 "Aseptic technique; vascular anatomy competence; emergency kit: hyaluronidase, anaphylaxis protocol, referral pathway.",
 "Adverse-event log; report to registries and authorities.",
 "Objective monitoring: cutometry, dermal ultrasound, standardized photography, trichoscopy/phototrichogram.",
])

H("8. Pretest & Posttest Answer Keys with Rationales")
table(["Pretest","Key","Rationale"], [
 ["Q1","b","Collagen I is ~80–85% of dermal collagen — the dominant structural protein"],
 ["Q2","a","PDRN's repair effects are mediated via adenosine A2A receptor activation"],
 ["Q3","b","Exosomes are 30–150 nm endosomal extracellular vesicles"],
 ["Q4","b","GHK-Cu is a copper-binding carrier tripeptide"],
 ["Q5","a","Senescent fibroblasts degrade ECM through SASP (IL-6/IL-8/MMPs)"],
 ["Q6","b","Chemical intelligence = instructive bioactive signaling molecules"],
 ["Q7","c","Injectable exosomes are unapproved in most jurisdictions"],
 ["Q8","b","Polynucleotides are polymers of nucleotides (DNA fragments)"],
 ["Q9","b","Palmitoyl pentapeptide mimics the KTTKS collagen I fragment"],
 ["Q10","b","Trichoscopy/phototrichogram is the standard objective hair measure"],
])
table(["Posttest","Key","Rationale"], [
 ["Q1","a","UV → AP-1 → MMP-1/3/9 drives collagen breakdown"],
 ["Q2","b","A2A activation is anti-inflammatory and pro-angiogenic"],
 ["Q3","b","Polynucleotides hold the strongest human RCT evidence (Level B)"],
 ["Q4","a","CD9/CD63/CD81 tetraspanins characterize exosomes"],
 ["Q5","b","Standard PN course: 3–4 sessions at 2–4 week intervals"],
 ["Q6","b","DP ECM volume correlates with follicle size and hair caliber"],
 ["Q7","b","Matrikines are ECM-derived signaling peptide fragments"],
 ["Q8","c","GHK-Cu declines ~60% between ages 20 and 60"],
 ["Q9","b","Injectable exosomes are investigational/unapproved in most regions"],
 ["Q10","b","PN is the evidence-graded first-line injectable for skin quality"],
])

H("9. Selected References (2015–2026)")
for r in [
 "Fisher GJ, et al. Mechanisms of photoaging, collagen fragmentation and fibroblast mechanosensing in aged skin. Am J Pathol / Arch Dermatol.",
 "Squadrito F, et al. PDRN and the adenosine A2A receptor in wound healing and tissue repair. Wound Repair Regen; Front Pharmacol.",
 "Cavallini M, et al. Polynucleotides in facial rejuvenation: split-face randomized data. J Cosmet Dermatol.",
 "Khan A, et al. Polynucleotide skin boosters: clinical and ultrasound outcomes. J Cosmet Dermatol.",
 "Gentile P, Garcovich S. Systematic reviews of extracellular vesicles/exosomes in skin rejuvenation and androgenetic alopecia. J Clin Med.",
 "Pickart L, Margolina A. Regenerative and protective actions of GHK-Cu. Int J Mol Sci, 2018.",
 "Fields K, et al. Topical signal peptides in photoaging: randomized controlled data. Int J Cosmet Sci.",
 "FDA. Public safety notification on exosome products (2019) and subsequent warnings on unapproved regenerative products.",
 "Narayan RS, et al. Hair follicle aging and dermal papilla ECM dysfunction. Exp Dermatol.",
 "3D Rejuvenation Academy CME module materials, 2026. www.3drejuvenationcode.com",
]:
    doc.add_paragraph(r, style="List Number")

doc.add_paragraph()
P("© 2026 Dr. Amr Ismail, MD | 3D Rejuvenation Academy — Approved CPD Provider | All rights reserved.", italic=True)
doc.save("/workspace/output/Expanded_Practitioner_Pocket_Guide.docx")
print("Pocket guide saved")

